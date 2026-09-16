"""数据层：读取 en.xlsx、定位 mp3、英文/中文断句与句级对齐。"""

from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from . import config

# --------------------------------------------------------------------------- 数据结构


@dataclass
class Sentence:
    """一句英文 + 对应的一句中文（对齐后）。"""

    index: int          # 全局序号，从 0 开始
    para: int           # 所属段落序号（用于第二遍的段落间距）
    en: str
    cn: str

    @property
    def words(self) -> list[str]:
        return self.en.split()


@dataclass
class Row:
    """en.xlsx 中的一行。"""

    excel_row: int              # Excel 行号（1 基，含表头行）
    level: str                  # A1 / A2 / B1 / B2
    title_en: str
    script_en: str
    mp3_name: str               # D 列原始文件名
    title_cn: str
    script_cn: str
    mp3_path: Path | None = None
    mp3_fallback: bool = False          # True = 靠"编号兜底"匹配上的，文件名并不一致
    sentences: list[Sentence] = field(default_factory=list)

    @property
    def id(self) -> str:
        return f"row{self.excel_row}"

    @property
    def stem(self) -> str:
        if self.mp3_path is not None:
            return self.mp3_path.stem
        return f"{self.level}-row{self.excel_row}"

    @property
    def has_cn(self) -> bool:
        return bool(self.script_cn.strip())

    def summary(self) -> str:
        if self.mp3_path is None:
            audio = "（缺音频）"
        elif self.mp3_fallback:
            audio = f"（文件名待确认）{self.mp3_path.name}"
        else:
            audio = self.mp3_path.name
        cn = "有中文" if self.has_cn else "无中文"
        return (f"row{self.excel_row:>4} {self.level:<3} {cn} "
                f"句数={len(self.sentences):>3} 音频={audio}")


# --------------------------------------------------------------------------- 读取表格


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.replace("\r\n", "\n").replace("\r", "\n").strip()
    return str(value).strip()


def read_rows(xlsx_path: Path | None = None) -> list[Row]:
    """读取 en.xlsx 第 1 个工作表，返回全部数据行（跳过表头）。"""
    path = Path(xlsx_path or config.XLSX_PATH)
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb.worksheets[0]

    rows: list[Row] = []
    for r, values in enumerate(ws.iter_rows(values_only=True), start=1):
        if r == 1:
            continue
        values = list(values) + [None] * (6 - len(values))
        level = _cell(values[0])
        title_en = _cell(values[1])
        script_en = _cell(values[2])
        mp3_name = _cell(values[3])
        title_cn = _cell(values[4])
        script_cn = _cell(values[5])
        if not any((level, title_en, script_en, mp3_name, title_cn, script_cn)):
            continue
        rows.append(Row(
            excel_row=r,
            level=level.upper(),
            title_en=title_en,
            script_en=script_en,
            mp3_name=mp3_name,
            title_cn=title_cn,
            script_cn=script_cn,
        ))
    wb.close()
    return rows


# --------------------------------------------------------------------------- 定位音频


def _norm(name: str) -> str:
    """文件名归一化：小写、去扩展名、非字母数字统一成下划线。"""
    stem = Path(name).stem.lower()
    stem = unicodedata.normalize("NFKD", stem)
    return re.sub(r"[^a-z0-9]+", "_", stem).strip("_")


def index_audio(data_dir: Path | None = None) -> dict[str, Path]:
    """把 data 目录里的音频建成索引： 归一化名 -> 路径。"""
    d = Path(data_dir or config.DATA_DIR)
    idx: dict[str, Path] = {}
    for p in sorted(d.glob("*.mp3")):
        idx[_norm(p.name)] = p
    return idx


_NUM_PREFIX = re.compile(r"^[a-z]\d[-_]\d{3}[-_]?")


def _desc_tokens(name: str) -> set[str]:
    """文件名里"编号之后"的描述性词元，用于判断两个文件名是不是同一段内容。"""
    stem = _NUM_PREFIX.sub("", Path(name).stem.lower())
    return {t for t in re.findall(r"[a-z][a-z0-9]*", stem)}


def resolve_audio(rows: Iterable[Row], data_dir: Path | None = None) -> None:
    """为每行就地匹配 mp3。

    匹配顺序：文件名精确 → 归一化（大小写/`_`/`-` 差异）→ 编号兜底。
    编号兜底只在"描述性词元有交集"时才生效，避免把 A1-014-ANNA-COOK.mp3
    错配到同编号但内容完全不同的 A1-014-JULIE-FOODS-DISLIKE.mp3。
    """
    idx = index_audio(data_dir)
    by_num: dict[str, list[Path]] = {}
    for key, p in idx.items():
        m = re.match(r"([a-z]\d)_(\d{3})", key)
        if m:
            by_num.setdefault(f"{m.group(1)}{m.group(2)}", []).append(p)

    for row in rows:
        row.mp3_path = None
        row.mp3_fallback = False
        if not row.mp3_name:
            continue
        if row.mp3_name in idx:
            row.mp3_path = idx[row.mp3_name]
            continue
        if _norm(row.mp3_name) in idx:
            row.mp3_path = idx[_norm(row.mp3_name)]
            continue

        m = re.search(r"([A-Za-z]\d)-(\d{3})", row.mp3_name)
        if not m:
            continue
        want = _desc_tokens(row.mp3_name)
        for cand in by_num.get(f"{m.group(1).lower()}{m.group(2)}", []):
            if want and (want & _desc_tokens(cand.name)):
                row.mp3_path = cand
                row.mp3_fallback = True
                break


# --------------------------------------------------------------------------- 断句

# 只有这些"确实是缩写"的形式才阻止断句；必须区分大小写，
# 否则 "us." / "no." / "am." 之类的普通词会被误判成缩写。
_ABBREV_TOKENS = {
    "Mr.", "Mrs.", "Ms.", "Dr.", "St.", "Sr.", "Jr.", "Prof.", "vs.", "etc.",
    "e.g.", "i.e.", "Inc.", "Ltd.", "Co.", "approx.", "Dept.", "Fig.",
    "U.S.", "U.K.", "U.N.", "a.m.", "p.m.", "Ph.D.", "Mt.", "Ave.",
}
_INITIAL = re.compile(r"^[A-Z]\.$")

_SENT_END_EN = re.compile(r"(?<=[.!?])[\"'\u201d\u2019)]*\s+")

_END_EN = ".!?\u2026"
_END_CN = "\u3002\uff01\uff1f\u2026"
_TRAILING = "\"'\u201d\u2019)]\uff09\u3011\u300d\u300f"


def _ends_sentence(line: str, enders: str) -> bool:
    s = line.rstrip()
    while s and s[-1] in _TRAILING:
        s = s[:-1].rstrip()
    return bool(s) and s[-1] in enders


def _merge_broken_lines(text: str, enders: str, joiner: str) -> list[str]:
    """把"断在句子中间"的换行合并回上一段。

    原始文稿里有 27 行是硬换行（例如 "... but I love" / "that song."），
    如果直接按换行分段，视频里就会出现只写着 "but I love" 的一页。
    这里只在"上一段没有以句末标点结束"时才合并，因此正常的段落划分不受影响；
    英文和中文用同一套规则，段落数依然保持一一对应。
    """
    paras: list[str] = []
    for raw in text.split("\n"):
        line = raw.strip()
        if not line:
            continue
        if paras and not _ends_sentence(paras[-1], enders):
            paras[-1] = paras[-1] + joiner + line
        else:
            paras.append(line)
    return paras


def split_en_paragraphs(text: str) -> list[list[str]]:
    """英文先按换行分段（断在句中的硬换行会被合并），再按句末标点分句。"""
    out: list[list[str]] = []
    for para in _merge_broken_lines(text, _END_EN, " "):
        parts: list[str] = []
        start = 0
        for m in _SENT_END_EN.finditer(para):
            head = para[start:m.end()].strip()
            token = re.split(r"\s+", head)[-1] if head else ""
            if token and (token in _ABBREV_TOKENS or _INITIAL.match(token)):
                continue
            if head:
                parts.append(head)
            start = m.end()
        tail = para[start:].strip()
        if tail:
            parts.append(tail)
        if parts:
            out.append(parts)
    return out


def split_cn_paragraphs(text: str) -> list[list[str]]:
    """中文先按换行分段（规则与英文一致），再按中文句末标点分句。"""
    out: list[list[str]] = []
    for para in _merge_broken_lines(text, _END_CN, ""):
        parts = [p.strip() for p in re.split(r"(?<=[。！？!?…])", para) if p.strip()]
        if parts:
            out.append(parts)
    return out


# --------------------------------------------------------------------------- 句级对齐


def _en_len(s: str) -> int:
    return len(re.findall(r"[A-Za-z0-9']+", s))


def _cn_len(s: str) -> int:
    return len(re.sub(r"\s+", "", s))


# 非 1:1 组合的惩罚：调大可以让对齐在句数不等时也尽量贴着 1:1
MOVE_PENALTY = 0.7
# 中文句里带逗号时，更可能本来就是由两个英文短句合并来的 —— 给一点折扣
COMMA_BONUS = 0.35
# 某一句在另一种语言里没有对应译句时的代价（例如英文的 "Yes." 中文没单独译）
SKIP_PENALTY = 1.8
# 长度比例偏离的截断上限：太小会让短句失去区分度，太大会被个别离群句带偏
RATIO_CAP = 2.5
_CN_COMMA = "，,、"


def _group_cost(en: list[str], cn: list[str], ratio: float) -> float:
    import math

    if not en or not cn:                     # 一边为空：允许"漏译"，但代价较高
        return SKIP_PENALTY
    a = sum(_en_len(s) for s in en)
    b = sum(_cn_len(s) for s in cn)
    if a == 0 or b == 0:
        return 1e6
    r = a / b
    # 与全局比例的偏离（对数域，做上限截断以抑制离群句）
    cost = min(abs(math.log(r / ratio)), RATIO_CAP)
    # 强烈偏好 1:1：译文基本是逐句对应的，合/拆只应在句数不等时发生
    cost += MOVE_PENALTY * (len(en) + len(cn) - 2)
    # 多个英文短句并成 1 个中文句时，中文里通常留有逗号
    if len(cn) == 1 and len(en) > 1 and any(c in cn[0] for c in _CN_COMMA):
        cost -= COMMA_BONUS
    return cost


def align_paragraph(en_sents: list[str], cn_sents: list[str]) -> list[tuple[list[str], list[str]]]:
    """把一个段落里的英文句与中文句做单调对齐。

    译文基本是逐句对应的，因此：
      * 句数相等 → 直接 1:1（这是绝大多数情况，避免长度启发式造成的错位）；
      * 句数不等 → 动态规划，在必须发生的合并/拆分之外一律保持 1:1。
    返回若干分组 [(英文句子列表, 中文句子列表), ...]。
    """
    m, n = len(en_sents), len(cn_sents)
    if m == 0 or n == 0:
        return []
    if m == n:
        return [([en_sents[i]], [cn_sents[i]]) for i in range(m)]

    total_en = sum(_en_len(s) for s in en_sents)
    total_cn = sum(_cn_len(s) for s in cn_sents) or 1
    ratio = max(total_en, 1) / total_cn

    INF = float("inf")
    dp = [[INF] * (n + 1) for _ in range(m + 1)]
    back: list[list[tuple[int, int] | None]] = [[None] * (n + 1) for _ in range(m + 1)]
    dp[0][0] = 0.0

    # 译文常把 2~3 个英文短句并成一句中文，也会把一句英文拆成两句中文；
    # (1,0)/(0,1) 用于处理"某一句没有对应译句"的情况（如英文的 Yes. / No.）
    moves = ((1, 1), (1, 2), (2, 1), (1, 3), (3, 1), (2, 2), (1, 0), (0, 1))
    for i in range(m + 1):
        for j in range(n + 1):
            cur = dp[i][j]
            if cur == INF:
                continue
            for di, dj in moves:
                ni, nj = i + di, j + dj
                if ni > m or nj > n:
                    continue
                cost = cur + _group_cost(en_sents[i:ni], cn_sents[j:nj], ratio)
                if cost < dp[ni][nj]:
                    dp[ni][nj] = cost
                    back[ni][nj] = (i, j)

    groups: list[tuple[list[str], list[str]]] = []
    i, j = m, n
    while i > 0 or j > 0:
        prev = back[i][j]
        if prev is None:                     # 理论上不会发生，兜底
            groups.append((en_sents[:i], cn_sents[:j]))
            break
        pi, pj = prev
        groups.append((en_sents[pi:i], cn_sents[pj:j]))
        i, j = pi, pj
    groups.reverse()
    return groups


def _split_cn_proportional(text: str, weights: list[int]) -> list[str]:
    """把一段中文按权重切成 len(weights) 份，尽量落在标点之后。"""
    if len(weights) <= 1:
        return [text]
    total = sum(weights) or 1
    n = len(text)
    marks = "。！？；，、：!?;,"
    cuts: list[int] = []
    acc = 0
    for w in weights[:-1]:
        acc += w
        target = int(round(n * acc / total))
        lo, hi = max(1, target - max(4, n // 12)), min(n - 1, target + max(4, n // 12))
        best, best_d = None, None
        for k in range(lo, hi + 1):
            if text[k - 1] in marks:
                d = abs(k - target)
                if best_d is None or d < best_d:
                    best, best_d = k, d
        cuts.append(best if best is not None else min(max(target, 1), n - 1))

    ordered, last = [], 0
    for c in cuts:
        c = max(c, last + 1)
        ordered.append(c)
        last = c
    pieces, prev = [], 0
    for c in ordered:
        pieces.append(text[prev:c].strip())
        prev = c
    pieces.append(text[prev:].strip())
    # 拆在逗号/顿号后面时，去掉悬挂的标点，字幕更干净
    pieces = [p[:-1].rstrip() if p and p[-1] in "，、,;" else p for p in pieces]
    return pieces


def build_sentences(script_en: str, script_cn: str) -> list[Sentence]:
    """把英文稿与中文稿对齐成逐句列表。

    先做**段落级**对齐（容忍英中段落数不一致），再在每组内做**句级**对齐。
    这样即使译文把两段并成一段、或把一段拆成两段，也能得到干净的逐句对照。
    中文缺失时只产出英文句。
    """
    en_paras = split_en_paragraphs(script_en)
    cn_paras = split_cn_paragraphs(script_cn) if script_cn.strip() else []

    sentences: list[Sentence] = []
    if not cn_paras:
        for pi, en_sents in enumerate(en_paras):
            for s in en_sents:
                sentences.append(Sentence(len(sentences), pi, s, ""))
        return sentences

    para_groups = align_paragraph([" ".join(p) for p in en_paras],
                                  ["".join(p) for p in cn_paras])

    ei = ci = 0
    for gi, (en_grp, cn_grp) in enumerate(para_groups):
        en_sents = [s for p in en_paras[ei:ei + len(en_grp)] for s in p]
        cn_sents = [s for p in cn_paras[ci:ci + len(cn_grp)] for s in p]
        ei += len(en_grp)
        ci += len(cn_grp)
        if not en_sents:
            continue
        if not cn_sents:
            for s in en_sents:
                sentences.append(Sentence(len(sentences), gi, s, ""))
            continue

        for en_g, cn_g in align_paragraph(en_sents, cn_sents):
            if not en_g:                                  # 中文多出来的一句，丢弃
                continue
            cn_text = "".join(cn_g).strip()
            if len(en_g) == 1:
                sentences.append(Sentence(len(sentences), gi, en_g[0], cn_text))
                continue
            # 多个英文短句对应更少的中文句：优先按标点把中文切成同样份数；
            # 切出来有太短的碎片时，退化为整句重复显示，避免把中文切得不成句。
            pieces = _split_cn_proportional(cn_text, [_en_len(s) for s in en_g])
            usable = (len(pieces) == len(en_g)
                      and all(len(p.strip()) >= 4 for p in pieces))
            for k, s in enumerate(en_g):
                sentences.append(Sentence(len(sentences), gi, s,
                                          pieces[k] if usable else cn_text))

    _fill_missing_cn(sentences)
    return sentences


def _fill_missing_cn(sentences: list[Sentence]) -> None:
    """兜底：任何一句中文为空时，沿用最近的非空中文，保证每页都有中文字幕。"""
    last = ""
    for s in sentences:
        if s.cn.strip():
            last = s.cn
        elif last:
            s.cn = last
    nxt = ""
    for s in reversed(sentences):
        if s.cn.strip():
            nxt = s.cn
        elif nxt:
            s.cn = nxt

# --------------------------------------------------------------------------- 组装


def load_dataset(xlsx_path: Path | None = None, data_dir: Path | None = None) -> list[Row]:
    """读取表格、匹配音频、完成断句与英中句对齐。

    如果 cache/rowN.lines.json 存在，则用它覆盖自动对齐结果（便于人工修订）。
    """
    rows = read_rows(xlsx_path)
    resolve_audio(rows, data_dir)
    for row in rows:
        row.sentences = build_sentences(row.script_en, row.script_cn)
        apply_lines_override(row)
    return rows


# --------------------------------------------------------------------------- 人工修订


def lines_path(row: Row) -> Path:
    return config.CACHE_DIR / f"{row.id}.lines.json"


def export_lines(row: Row) -> Path:
    """把自动对齐结果导出成可编辑文件，用于人工核对 / 修订。"""
    import json

    config.CACHE_DIR.mkdir(parents=True, exist_ok=True)
    path = lines_path(row)
    path.write_text(json.dumps({
        "title_en": row.title_en,
        "title_cn": row.title_cn,
        "_说明": "可直接修改每条的 en / cn / para；改完重新运行 make_video.py 即生效。"
                 "删除本文件即回到自动对齐。",
        "lines": [{"para": s.para, "en": s.en, "cn": s.cn} for s in row.sentences],
    }, ensure_ascii=False, indent=2), "utf-8")
    return path


def apply_lines_override(row: Row) -> bool:
    """若存在人工修订文件，则覆盖自动对齐结果。"""
    import json

    path = lines_path(row)
    if not path.exists():
        return False
    try:
        data = json.loads(path.read_text("utf-8"))
    except Exception:
        return False
    lines = data.get("lines") or []
    if not lines:
        return False
    row.sentences = [
        Sentence(index=i, para=int(item.get("para", 0)),
                 en=str(item.get("en", "")).strip(), cn=str(item.get("cn", "")).strip())
        for i, item in enumerate(lines)
    ]
    return True


def select_rows(rows: list[Row], spec: str | None = None, level: str | None = None,
                only_with_audio: bool = False, only_with_cn: bool = False) -> list[Row]:
    """按 --rows / --level 过滤。spec 支持 "2,3,5-8" 与 mp3 文件名片段。"""
    picked = rows
    if level:
        want = {x.strip().upper() for x in level.split(",") if x.strip()}
        picked = [r for r in picked if r.level in want]
    if spec:
        want_rows: set[int] = set()
        want_name: list[str] = []
        for token in spec.split(","):
            token = token.strip()
            if not token:
                continue
            m = re.fullmatch(r"(\d+)\s*-\s*(\d+)", token)
            if m:
                a, b = int(m.group(1)), int(m.group(2))
                want_rows.update(range(min(a, b), max(a, b) + 1))
            elif token.isdigit():
                want_rows.add(int(token))
            else:
                want_name.append(token.lower())
        if want_rows or want_name:
            picked = [r for r in picked
                      if r.excel_row in want_rows
                      or any(t in (r.mp3_name or "").lower() or t in r.stem.lower()
                             for t in want_name)]
    if only_with_audio:
        picked = [r for r in picked if r.mp3_path]
    if only_with_cn:
        picked = [r for r in picked if r.has_cn]
    return picked
