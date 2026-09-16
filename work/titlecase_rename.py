"""一次性脚本：把 en.xlsx 的 mp3 列 / data/ / output/ 文件名改成首字母大写。

规则：每个英文单词首字母大写、其余小写；短虚词（冠词/介词/连词/be 动词）
在非首尾位置保持小写。仅改字母大小写，分隔符（- _ 空格）与扩展名保持不变。
"""
import json
import os
import re
import shutil
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "data" / "en.xlsx"

SMALL = {"a", "an", "the", "and", "or", "but", "nor", "for", "so", "yet",
         "as", "at", "by", "in", "into", "of", "on", "onto", "to", "up",
         "with", "from", "off", "out", "over",
         "is", "are", "was", "were", "be", "been", "am"}
SEP = "-_ "
TOKEN = re.compile(r"^([^A-Za-z]*)([A-Za-z][A-Za-z'’]*)([^A-Za-z]*)$")


def cap_word(word: str) -> str:
    m = TOKEN.match(word)
    if not m:
        return word
    pre, core, post = m.groups()
    return pre + core[0].upper() + core[1:].lower() + post


def titlecase(name: str) -> str:
    stem, dot, ext = name.rpartition(".")
    if not dot:
        stem, ext = name, ""
    parts = re.split(r"([\-_ ])", stem)
    words = [i for i, p in enumerate(parts) if p and p not in SEP]
    out = []
    for i, part in enumerate(parts):
        if not part or part in SEP:
            out.append(part)
            continue
        c = cap_word(part)
        if words and i not in (words[0], words[-1]) and c.lower() in SMALL:
            c = c.lower()
        out.append(c)
    return "".join(out) + (dot + ext if dot else "")


def plan(directory: Path, pattern: str) -> list[tuple[Path, Path]]:
    moves = []
    for path in sorted(directory.glob(pattern)):
        if not path.is_file():
            continue
        new = path.with_name(titlecase(path.name))
        if new != path:
            moves.append((path, new))
    return moves


def check_collisions(moves, label):
    targets = {}
    for old, new in moves:
        key = str(new).lower()
        if key in targets and targets[key] != old:
            sys.exit(f"[{label}] 目标名冲突：{old} / {targets[key]} -> {new}")
        targets[key] = old
    for old, new in moves:
        if new.exists() and new != old and new.name.lower() != old.name.lower():
            sys.exit(f"[{label}] 目标已存在：{new}")


def rename_all(moves):
    done = []
    for old, new in moves:
        if old.name.lower() == new.name.lower():
            tmp = old.with_name(old.name + ".renaming-tmp")
            os.rename(old, tmp)
            os.rename(tmp, new)
        else:
            os.rename(old, new)
        done.append((str(old), str(new)))
    return done


def main(apply: bool):
    wb = load_workbook(XLSX)
    ws = wb["Sheet1"]
    log = {"time": datetime.now().isoformat(timespec="seconds"), "applied": apply, "xlsx": [], "files": []}

    # 1) xlsx D 列
    for r in range(2, ws.max_row + 1):
        cell = ws.cell(row=r, column=4)
        if not cell.value:
            continue
        new = titlecase(str(cell.value))
        if new != cell.value:
            log["xlsx"].append({"row": r, "old": cell.value, "new": new})
            if apply:
                cell.value = new
    print(f"xlsx D 列需改 {len(log['xlsx'])} 项")

    # 2) data/ output/ 媒体文件（含子目录）
    moves = plan(ROOT / "data", "**/*.mp3") + plan(ROOT / "output", "**/*.mp4")
    check_collisions(moves, "media")
    print(f"媒体文件需改 {len(moves)} 项")

    # 3) 按 mp3 文件名做 key 的缓存
    cache_moves = plan(ROOT / "cache", "*.timings.json") \
        + plan(ROOT / "cache" / "backup-silence", "*.json") \
        + plan(ROOT / "cache" / "asr-cache", "*.json")
    check_collisions(cache_moves, "cache")
    print(f"缓存文件需改 {len(cache_moves)} 项")

    if not apply:
        for old, new in moves[:8]:
            print(f"  {old.relative_to(ROOT)}  ->  {new.name}")
        return

    shutil.copy2(XLSX, ROOT / "data" / f"en.backup-{datetime.now():%Y%m%d-%H%M%S}.xlsx")
    wb.save(XLSX)
    log["files"] = rename_all(moves) + rename_all(cache_moves)

    out = ROOT / "work" / f"titlecase-rename-{datetime.now():%Y%m%d-%H%M%S}.json"
    out.write_text(json.dumps(log, ensure_ascii=False, indent=2), "utf-8")
    print("日志：", out)


if __name__ == "__main__":
    main("--apply" in sys.argv)
