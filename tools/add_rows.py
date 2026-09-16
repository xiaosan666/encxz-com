#!/usr/bin/env python3
"""把 work/new_rows.json 里的新行追加进 en.xlsx，然后重新排序。

每行的字段：level / title_en / title_cn / mp3 / paras_en[] / paras_cn[]
paras_en 与 paras_cn 的段数必须一致（脚本会校验）。

    python tools/add_rows.py --file work/new_rows.json --dry-run
    python tools/add_rows.py --file work/new_rows.json
"""

from __future__ import annotations

import argparse
import json
import shutil
import sys
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import load_workbook                          # noqa: E402

from video_gen import config                                # noqa: E402
from tools.normalize_table import clean_title, sort_key     # noqa: E402


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--file", required=True)
    ap.add_argument("--xlsx", default=str(config.XLSX_PATH))
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    new = json.loads(Path(args.file).read_text("utf-8"))
    xlsx = Path(args.xlsx)
    wb = load_workbook(xlsx)
    ws = wb.worksheets[0]

    data = []
    for r in range(2, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, 7)]
        if not any(v not in (None, "") for v in vals):
            continue
        data.append({"lv": vals[0], "b": vals[1], "c": vals[2],
                     "d": vals[3], "e": vals[4], "f": vals[5]})
    before = len(data)
    have = {str(d["d"] or "") for d in data}
    have_num = {str(d["d"] or "")[:6] for d in data}

    problems = []
    for item in new:
        if item["mp3"] in have:
            problems.append(f"{item['mp3']} 已存在，跳过")
            continue
        if str(item["mp3"])[:6] in have_num:
            problems.append(f"{item['mp3'][:6]} 编号已被占用：{item['mp3']}")
            continue
        if len(item["paras_en"]) != len(item["paras_cn"]):
            problems.append(f"{item['mp3']} 英中段数不一致")
            continue
        data.append({
            "lv": item["level"],
            "b": clean_title(item["title_en"]),
            "c": "\n".join(item["paras_en"]),
            "d": item["mp3"],
            "e": clean_title(item["title_cn"]),
            "f": "\n".join(item["paras_cn"]),
        })

    if problems:
        print("有问题，未写回：")
        for p in problems:
            print("  ·", p)
        return 2

    ordered = sorted(data, key=sort_key)
    print(f"{before} 行 + 新增 {len(new)} 行 = {len(ordered)} 行")
    for item in new:
        print(f"  + {item['level']}  {item['mp3']:<44} {item['title_en'][:44]}")

    if args.dry_run:
        print("\n--dry-run：未写回。")
        return 0

    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    backup = xlsx.with_name(f"{xlsx.stem}.backup-{stamp}{xlsx.suffix}")
    shutil.copy2(xlsx, backup)
    print(f"\n已备份 → {backup.name}")

    for i, row in enumerate(ordered, start=2):
        for c, k in enumerate(("lv", "b", "c", "d", "e", "f"), start=1):
            ws.cell(i, c).value = row[k]
    for r in range(len(ordered) + 2, ws.max_row + 1):
        for c in range(1, 7):
            ws.cell(r, c).value = None
    wb.save(xlsx)
    print(f"已写回 {len(ordered)} 行（已重新排序）")
    return 0


if __name__ == "__main__":
    sys.exit(main())
