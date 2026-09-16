#!/usr/bin/env python3
"""整理 en.xlsx：清理标题前缀 + 按等级与 mp3 编号排序。

1. 标题清理：去掉 `Beginner English - 11 -` / `初级英语 - 11 -` 这类系列前缀，
   只保留问题本身（B 列英文标题、E 列中文标题都处理）；
2. 排序：先按等级 A1 → A2 → B1 → B2 → C1，再按 mp3 文件名里的编号，
   最后按标题（保证结果稳定可复现）。

改前自动备份。用法：
    python tools/normalize_table.py --dry-run
    python tools/normalize_table.py
"""

from __future__ import annotations

import argparse
import re
import shutil
import sys
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import load_workbook                          # noqa: E402

from video_gen import config                                # noqa: E402

LEVEL_ORDER = {"A1": 0, "A2": 1, "B1": 2, "B2": 3, "C1": 4, "C2": 5}

EN_PREFIX = re.compile(
    r"^\s*(?:Beginner|Intermediate|High-Intermediate|Advanced|Elementary)\s+English"
    r"\s*[-\u2013\u2014]\s*\d+\s*[-\u2013\u2014]\s*", re.IGNORECASE)
CN_PREFIX = re.compile(
    r"^\s*(?:初级英语|中级英语|中高级英语|高级英语|入门英语|初级|中级|高级)"
    r"\s*[-\u2013\u2014]\s*\d+\s*[-\u2013\u2014]\s*")
NUM_IN_MP3 = re.compile(r"^[A-Za-z]\d-(\d{3})")


def clean_title(t: str) -> str:
    t = EN_PREFIX.sub("", t or "").strip()
    t = CN_PREFIX.sub("", t or "").strip()
    return t


def sort_key(row: dict):
    m = NUM_IN_MP3.match(row["d"] or "")
    num = int(m.group(1)) if m else 9999
    return (LEVEL_ORDER.get((row["lv"] or "").upper(), 9), num,
            (row["b"] or "").lower())


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", default=str(config.XLSX_PATH))
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()
    xlsx = Path(args.xlsx)

    wb = load_workbook(xlsx)
    ws = wb.worksheets[0]
    header = [c.value for c in ws[1]]

    data = []
    for r in range(2, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, 7)]
        if not any(v not in (None, "") for v in vals):
            continue
        data.append({"lv": (vals[0] or "").strip() if isinstance(vals[0], str) else vals[0],
                     "b": vals[1], "c": vals[2], "d": vals[3], "e": vals[4], "f": vals[5]})
    print(f"读入 {len(data)} 行，表头：{header}")

    changed = 0
    for row in data:
        for key in ("b", "e"):
            old = row[key] or ""
            new = clean_title(old)
            if new != old:
                changed += 1
                row[key] = new
    print(f"标题前缀已清理：{changed} 处")

    ordered = sorted(data, key=sort_key)
    moved = sum(1 for a, b in zip(data, ordered) if a is not b)
    print(f"排序后位置发生变化的行：{moved}")

    order = [f"{r['lv']}-{(NUM_IN_MP3.match(r['d'] or '').group(1) if NUM_IN_MP3.match(r['d'] or '') else '???')}"
             for r in ordered]
    print("排序结果预览（前 8 / 后 4）：")
    print("   " + "  ".join(order[:8]) + " ... " + "  ".join(order[-4:]))

    if args.dry_run:
        print("\n--dry-run：未写回。")
        return 0

    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    backup = xlsx.with_name(f"{xlsx.stem}.backup-{stamp}{xlsx.suffix}")
    shutil.copy2(xlsx, backup)
    print(f"\n已备份原表 → {backup.name}")

    for i, row in enumerate(ordered, start=2):
        for c, key in enumerate(("lv", "b", "c", "d", "e", "f"), start=1):
            ws.cell(i, c).value = row[key]
    # 清掉多余的历史行
    for r in range(len(ordered) + 2, ws.max_row + 1):
        for c in range(1, 7):
            ws.cell(r, c).value = None
    wb.save(xlsx)
    print(f"已写回 {len(ordered)} 行")
    return 0


if __name__ == "__main__":
    sys.exit(main())
