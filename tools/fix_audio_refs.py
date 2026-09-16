#!/usr/bin/env python3
"""修正 en.xlsx 里 D 列（mp3 文件名）的引用。

用途：音频文件其实存在于 data/，但表格 D 列指向了别的文件名 —— 本工具按
给定的「行号 → 文件名」映射改回来，改前自动备份，并校验文件确实存在。

用法
----
    python tools/fix_audio_refs.py --map 55=A2-005-WHAT-DO-YOU-MISS-ABOUT-YOUR-COUNTRY.mp3 \
                                        82=A2-032-WHEN-DO-YOU-DRINK-YOUR-FAVORITE-DRINK.mp3 \
                                        ... --dry-run
    python tools/fix_audio_refs.py --map ...          # 真正写回
"""

from __future__ import annotations

import argparse
import sys
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import load_workbook                          # noqa: E402

from video_gen import config                                # noqa: E402
from video_gen.data import index_audio, load_dataset        # noqa: E402


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--map", nargs="+", required=True,
                    help="形如 82=A2-032-xxx.mp3 的映射，可写多个")
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--xlsx", default=str(config.XLSX_PATH))
    args = ap.parse_args()

    xlsx = Path(args.xlsx)
    avail = {p.name for p in index_audio().values()}
    rows = {r.excel_row: r for r in load_dataset()}

    plan: list[tuple[int, str, str]] = []
    problems: list[str] = []
    for item in args.map:
        if "=" not in item:
            problems.append(f"格式不对：{item}")
            continue
        rn_s, fname = item.split("=", 1)
        rn_s, fname = rn_s.strip(), fname.strip()
        if not rn_s.isdigit():
            problems.append(f"行号不是数字：{rn_s}")
            continue
        rn = int(rn_s)
        if rn not in rows:
            problems.append(f"row{rn} 在表格里不存在")
            continue
        if fname not in avail:
            problems.append(f"row{rn} 指定的文件在 data/ 里找不到：{fname}")
            continue
        plan.append((rn, rows[rn].mp3_name, fname))

    if problems:
        print(f"发现 {len(problems)} 个问题：")
        for p in problems:
            print("  ·", p)
        if not args.dry_run:
            print("\n有问题时不写回。")
            return 2

    print(f"{'行':>5}  {'原 D 列':<50} → 新 D 列")
    for rn, old, new in plan:
        mark = "（不变）" if old == new else ""
        print(f"{rn:>5}  {old[:48]:<50} → {new} {mark}")

    if args.dry_run:
        print("\n--dry-run：未写回。")
        return 0

    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    backup = xlsx.with_name(f"{xlsx.stem}.backup-{stamp}{xlsx.suffix}")
    import shutil
    shutil.copy2(xlsx, backup)
    print(f"\n已备份原表 → {backup.name}")

    wb = load_workbook(xlsx)
    ws = wb.worksheets[0]
    for rn, _old, new in plan:
        ws.cell(rn, 4).value = new
    wb.save(xlsx)
    print(f"已更新 {len(plan)} 行的 D 列")
    return 0


if __name__ == "__main__":
    sys.exit(main())
