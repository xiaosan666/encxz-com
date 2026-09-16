#!/usr/bin/env python3
"""把 work/translations/*.json 里的中文译文校验后写回 data/en.xlsx。

- 校验：行号存在、paras_cn 段数与英文一致、无空值、无遗漏
- 写回：E 列（Title 中文）、F 列（script 中文，段落用换行拼接）
- 备份：写回前自动把原文件复制成 data/en.backup-<时间戳>.xlsx
- 已有中文的行不会被覆盖

用法
----
    python tools/apply_translations.py --dry-run   # 只校验，不写
    python tools/apply_translations.py             # 校验并写回
"""

from __future__ import annotations

import argparse
import json
import shutil
import sys
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import load_workbook                          # noqa: E402

from video_gen import config                                # noqa: E402
from video_gen.data import load_dataset                     # noqa: E402

WORK = config.ROOT / "work"


def collect(paths: list[Path]) -> tuple[dict[int, dict], list[str]]:
    """读取并合并所有批次文件，返回 (row -> 译文, 问题列表)。"""
    merged: dict[int, dict] = {}
    problems: list[str] = []
    for p in sorted(paths):
        try:
            data = json.loads(p.read_text("utf-8"))
        except Exception as exc:                            # noqa: BLE001
            problems.append(f"{p.name}: JSON 解析失败 {exc}")
            continue
        if not isinstance(data, list):
            problems.append(f"{p.name}: 顶层不是数组")
            continue
        for item in data:
            row = item.get("row")
            if row in merged:
                problems.append(f"{p.name}: row{row} 重复出现")
            merged[row] = item
    return merged, problems


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true", help="只校验不写回")
    ap.add_argument("--xlsx", default=str(config.XLSX_PATH))
    ap.add_argument("--dir", default=str(WORK / "translations"))
    args = ap.parse_args()

    xlsx = Path(args.xlsx)
    files = sorted(Path(args.dir).glob("batch_*.json"))
    if not files:
        print(f"{args.dir} 下没有 batch_*.json")
        return 1
    print(f"读取 {len(files)} 个批次文件")

    merged, problems = collect(files)

    # ---------------- 校验 ----------------
    rows = {r.excel_row: r for r in load_dataset()}
    need = {rn for rn, r in rows.items() if not r.has_cn}

    missing = sorted(need - set(merged))
    extra = sorted(set(merged) - set(rows))
    if missing:
        problems.append(f"以下 {len(missing)} 行没有译文：{missing}")
    if extra:
        problems.append(f"译文里出现表格中不存在的行号：{extra}")

    ok: dict[int, dict] = {}
    for rn, item in sorted(merged.items()):
        row = rows.get(rn)
        if row is None:
            continue
        paras_en = [p.strip() for p in row.script_en.split("\n") if p.strip()]
        paras_cn = [str(p).strip() for p in (item.get("paras_cn") or [])]
        title_cn = str(item.get("title_cn", "")).strip()

        if len(paras_cn) != len(paras_en):
            problems.append(f"row{rn}: 段落数不一致 英{len(paras_en)} / 中{len(paras_cn)}")
            continue
        if any(not p for p in paras_cn):
            problems.append(f"row{rn}: 存在空段落")
            continue
        if not title_cn:
            problems.append(f"row{rn}: title_cn 为空")
            continue
        ok[rn] = {"title_cn": title_cn, "script_cn": "\n".join(paras_cn),
                  "paras": len(paras_cn)}

    total_paras = sum(v["paras"] for v in ok.values())
    print(f"校验通过 {len(ok)} 行（共 {total_paras} 段）；需要翻译 {len(need)} 行")

    if problems:
        print(f"\n发现 {len(problems)} 个问题：")
        for p in problems[:40]:
            print("  ·", p)
        if not args.dry_run:
            print("\n有问题时不写回，请先修正。")
            return 2
    if args.dry_run:
        print("\n--dry-run：未写回。")
        return 0

    # ---------------- 写回 ----------------
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    backup = xlsx.with_name(f"{xlsx.stem}.backup-{stamp}{xlsx.suffix}")
    shutil.copy2(xlsx, backup)
    print(f"\n已备份原表 → {backup.name}")

    wb = load_workbook(xlsx)
    ws = wb.worksheets[0]
    written = 0
    for rn, v in ok.items():
        if str(ws.cell(rn, 6).value or "").strip():
            continue                                   # 已有中文，不覆盖
        ws.cell(rn, 5).value = v["title_cn"]
        ws.cell(rn, 6).value = v["script_cn"]
        written += 1
    wb.save(xlsx)
    print(f"已写回 {written} 行（E 列中文标题 + F 列中文文稿）")
    return 0


if __name__ == "__main__":
    sys.exit(main())
